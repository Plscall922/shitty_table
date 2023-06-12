import os
from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.styles import PatternFill

class SaveTableSomewhere:

    def __init__(self, path, tables):
        self.path = path
        self.tables = tables
        self.savemytable()
    def savemytable(self):
        path = self.path
        tables = self.tables
        tables.save(path)

class GenerateFuckingTable:

    def __init__(self):
        self.execute
        self.table = []
        self.path = "D:/WORKSPACE/table.xlsx"


    def generatetable(self):

        self.table = Workbook()
        tables = self.table.active

        data = []

        dict = {
            'Apples_fruits_212': [10000, 5000, 8000, 6000],
            'Pears_fruits_212': [2000, 3000, 4000, 5000],
            'Bananas_fruits_212': [6000, 6000, 6500, 6000],
            'Oranges_fruits_212': [500, 300, 200, 700],
        }


        for key, values in dict.items():
            for index, value in enumerate(values):
                if index == 0:
                    data.append([key, value])
                else:
                    data.append(["", value])



        for row in data:
            tables.append(row)
            print(row)

        column_index = 'A'

        # Iterate over the cells in the column to determine the maximum content width
        max_width = 0
        for cell in tables[column_index]:
            if cell.value:
                cell_width = len(str(cell.value))
                if cell_width > max_width:
                    max_width = cell_width

        column_letter = column_index.upper()
        tables.column_dimensions[column_letter].width = max_width



    def execute(self):
        self.generatetable()
        path = self.path
        table = self.table
        SaveTableSomewhere(path, table)


hh = GenerateFuckingTable()

print(hh.execute())

